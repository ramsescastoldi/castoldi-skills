# -*- coding: utf-8 -*-
"""
parser_relatorio.py — Parser do relatório "Margem de Lucro de Combustíveis" (.xlsx)

Percorre um ou mais arquivos .xlsx no formato descrito em formato_relatorio.md
e devolve um DataFrame consolidado com as colunas exatas do Passo 4 do SKILL.md:

    Data, Filial, Produto, P_Custo, P_Venda, Margem_Bruta_Pct, Markup_Pct,
    Litragem, Total_Custo, Total_Venda, Lucro_Bruto

Regras implementadas:
- A filial é detectada pelo cabeçalho interno "Filial: XXX - NOME"
  (NUNCA pelo nome do arquivo). Formato "502 - POSTO 10 TRUCK" é mantido.
- Os produtos vêm de linhas "Produto: 000XXX - NOME" — o código é removido
  e o nome é normalizado ("000002 - ETANOL" -> "Etanol", siglas S10/S500 mantidas).
- Linhas de dados são reconhecidas por Dt. Venda ser datetime.
- Linhas "Subtotal" e "Total" são IGNORADAS (totais são sempre recalculados).
- Código defensivo: relatório fora do formato -> warning e o arquivo é pulado.

Uso:
    python3 parser_relatorio.py arquivo1.xlsx arquivo2.xlsx ...
ou, em código:
    from parser_relatorio import processar_arquivos
    df_consolidado = processar_arquivos(["a.xlsx", "b.xlsx"])
"""

import re
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Colunas exatas do DataFrame de saída (Passo 4 do SKILL.md)
COLUNAS_SAIDA = [
    "Data", "Filial", "Produto",
    "P_Custo", "P_Venda", "Margem_Bruta_Pct", "Markup_Pct",
    "Litragem", "Total_Custo", "Total_Venda", "Lucro_Bruto",
]

# Siglas que devem permanecer em maiúsculas ao normalizar nomes de produto
_SIGLAS = {"S10", "S500", "GNV", "B12", "B14"}

# Prefixos de linhas de totais que devem ser ignoradas
_LINHAS_IGNORAR = ("subtotal", "total")


def limpar_produto(texto):
    """'Produto: 000002 - ETANOL' -> 'Etanol' | '000009 - DIESEL S10' -> 'Diesel S10'."""
    nome = re.sub(r"^\s*produto\s*:\s*", "", str(texto), flags=re.IGNORECASE)
    nome = re.sub(r"^\s*\d+\s*-\s*", "", nome).strip()
    palavras = []
    for palavra in nome.split():
        if palavra.upper() in _SIGLAS:
            palavras.append(palavra.upper())
        else:
            palavras.append(palavra.capitalize())
    return " ".join(palavras)


def limpar_filial(texto):
    """'Filial: 502 - POSTO 10 TRUCK' -> '502 - POSTO 10 TRUCK' (formato mantido)."""
    nome = re.sub(r"^\s*filial\s*:\s*", "", str(texto), flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", nome).strip()


def _eh_texto(celula):
    return isinstance(celula, str) and celula.strip() != ""


def _eh_data(celula):
    return isinstance(celula, (datetime, date))


def _numero(celula, padrao=0.0):
    """Converte célula em float de forma defensiva (None/texto -> padrao)."""
    if isinstance(celula, (int, float)) and not isinstance(celula, bool):
        return float(celula)
    if isinstance(celula, str):
        try:
            return float(celula.replace(".", "").replace(",", "."))
        except ValueError:
            return padrao
    return padrao


def parsear_relatorio(caminho):
    """
    Lê UM arquivo .xlsx no formato "Margem de Lucro de Combustíveis" e devolve
    um DataFrame com as colunas COLUNAS_SAIDA.

    Se o arquivo estiver fora do formato (sem 'Filial:' ou sem linhas de dados),
    emite warning e devolve DataFrame vazio (o arquivo é pulado na consolidação).
    """
    caminho = Path(caminho)
    try:
        wb = load_workbook(caminho, data_only=True, read_only=True)
    except Exception as exc:  # arquivo corrompido, não é xlsx etc.
        warnings.warn(
            f"[{caminho.name}] não pôde ser aberto como .xlsx ({exc}). Arquivo pulado."
        )
        return pd.DataFrame(columns=COLUNAS_SAIDA)

    ws = wb.worksheets[0]  # relatório é sempre a primeira aba
    filial = None
    produto_atual = None
    registros = []

    for linha in ws.iter_rows(values_only=True):
        if not linha:
            continue
        primeira = linha[0]

        # --- Linha "Filial: XXX - NOME" (identifica a filial pelo cabeçalho interno)
        if _eh_texto(primeira) and re.match(r"^\s*filial\s*:", primeira, re.IGNORECASE):
            filial = limpar_filial(primeira)
            continue

        # --- Linha "Produto: 000XXX - NOME" (abre um novo bloco)
        if _eh_texto(primeira) and re.match(r"^\s*produto\s*:", primeira, re.IGNORECASE):
            produto_atual = limpar_produto(primeira)
            continue

        # --- Linhas Subtotal/Total: IGNORAR (recalculamos tudo na consolidação)
        if _eh_texto(primeira) and primeira.strip().lower().startswith(_LINHAS_IGNORAR):
            continue

        # --- Linha de dados: Dt. Venda vem como datetime nativo do Excel
        if _eh_data(primeira):
            if produto_atual is None:
                # dado sem bloco de produto — formato inesperado, pula a linha
                warnings.warn(
                    f"[{caminho.name}] linha de dados sem 'Produto:' anterior — linha pulada."
                )
                continue
            # Layout: Dt.Venda | Dias | P.Custo | P.Venda | Venda-Custo |
            #         % Mrg Bruta | % Markup | Litragem | T.Custo | T.Venda | L.Bruto
            if len(linha) < 11:
                warnings.warn(
                    f"[{caminho.name}] linha de dados com menos de 11 colunas — linha pulada."
                )
                continue
            dt = primeira.date() if isinstance(primeira, datetime) else primeira
            registros.append({
                "Data": dt,
                "Filial": filial,  # preenchida abaixo se a linha Filial vier depois
                "Produto": produto_atual,
                "P_Custo": _numero(linha[2]),
                "P_Venda": _numero(linha[3]),
                "Margem_Bruta_Pct": _numero(linha[5]),
                "Markup_Pct": _numero(linha[6]),
                "Litragem": _numero(linha[7]),
                "Total_Custo": _numero(linha[8]),
                "Total_Venda": _numero(linha[9]),
                "Lucro_Bruto": _numero(linha[10]),
            })
            continue

        # Demais linhas (título, período, cabeçalho de colunas, vazias): ignorar

    wb.close()

    # --- Validações defensivas: relatório fora do formato -> warning e pula
    if filial is None:
        warnings.warn(
            f"[{caminho.name}] cabeçalho 'Filial: XXX - NOME' não encontrado — "
            "relatório fora do formato esperado. Arquivo pulado."
        )
        return pd.DataFrame(columns=COLUNAS_SAIDA)

    if not registros:
        warnings.warn(
            f"[{caminho.name}] nenhuma linha de dados reconhecida — "
            "relatório fora do formato esperado. Arquivo pulado."
        )
        return pd.DataFrame(columns=COLUNAS_SAIDA)

    # Garante a filial em todas as linhas (caso a linha 'Filial:' venha após dados)
    df = pd.DataFrame(registros, columns=COLUNAS_SAIDA)
    df["Filial"] = df["Filial"].fillna(filial)
    return df


def processar_arquivos(caminhos):
    """
    Processa vários arquivos e concatena tudo em um único df_consolidado.
    Arquivos fora do formato geram warning e são pulados (nunca inventar dados).
    """
    frames = []
    for caminho in caminhos:
        df = parsear_relatorio(caminho)
        if df.empty:
            continue
        frames.append(df)
        print(
            f"OK  {Path(caminho).name}: {len(df)} registros | "
            f"filial {df['Filial'].iloc[0]} | {df['Produto'].nunique()} produtos"
        )

    if not frames:
        warnings.warn("Nenhum arquivo válido foi processado — df_consolidado vazio.")
        return pd.DataFrame(columns=COLUNAS_SAIDA)

    df_consolidado = pd.concat(frames, ignore_index=True)
    df_consolidado = df_consolidado.sort_values(
        ["Filial", "Produto", "Data"]
    ).reset_index(drop=True)
    return df_consolidado


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        print("Uso: python3 parser_relatorio.py arquivo1.xlsx [arquivo2.xlsx ...]")
        return 1

    df_consolidado = processar_arquivos(argv)
    if df_consolidado.empty:
        print("Nenhum dado consolidado (todos os arquivos foram pulados).")
        return 2

    print("\n=== df_consolidado ===")
    print(f"{len(df_consolidado)} registros | "
          f"{df_consolidado['Filial'].nunique()} filiais | "
          f"{df_consolidado['Produto'].nunique()} produtos | "
          f"período {df_consolidado['Data'].min()} a {df_consolidado['Data'].max()}")
    print(df_consolidado.head(10).to_string(index=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
